"""Writing the scorecard workbook: the template's formulas must survive, the
9 input cells must be replaced, and every figure must be traceable."""

import os

import pytest
from openpyxl import load_workbook

from compute_metrics import ExtractedFinancials, compute_metrics, render_metrics_for_excel
from fill_excel import (AUDIT_SHEET, METRIC_COLUMNS, METRIC_ROW, TARGET_SHEET,
                        find_template, read_scorecard_values,
                        write_metrics_to_template)

from test_compute_metrics import ZENCODER_FY2011

RATING_ROW = 6  # "ElCap Rating" formulas, directly under the input row


@pytest.fixture
def zencoder_output(tmp_path):
    fin = ExtractedFinancials.from_dict(ZENCODER_FY2011)
    metrics = compute_metrics(fin)
    out = str(tmp_path / "zencoder.xlsx")
    write_metrics_to_template(
        render_metrics_for_excel(metrics), out,
        financials=fin, metrics=metrics,
        source_file="Audited Financial Statements.pdf",
        warnings=fin.validate(), disagreements=["total revenue: 1 vs 2"],
    )
    return out, fin, metrics


def test_template_is_present():
    assert os.path.isfile(find_template())


def test_values_land_in_d5_to_l5(zencoder_output):
    out, _fin, metrics = zencoder_output
    written = read_scorecard_values(out)
    assert written[0] == 707868
    assert written[4] == pytest.approx(888160 / 248149, rel=1e-6)
    assert len(written) == len(METRIC_COLUMNS) == 9


def test_uncomputable_metrics_are_blank_not_stale(zencoder_output):
    """The shipped template has a worked example in D5:L5; none of it may
    survive into a generated workbook."""
    out, _fin, _metrics = zencoder_output
    written = read_scorecard_values(out)
    assert written[2] is None      # DSCR
    assert written[3] is None      # EBITDA / Interest
    template_values = read_scorecard_values(find_template())
    assert template_values[2] is not None   # the template really does hold one
    assert written[2] != template_values[2]


def test_rating_formulas_are_preserved(zencoder_output):
    out, _fin, _metrics = zencoder_output
    ws = load_workbook(out)[TARGET_SHEET]
    for col in METRIC_COLUMNS:
        value = ws[f"{col}{RATING_ROW}"].value
        text = getattr(value, "text", value)   # array formulas carry .text
        assert isinstance(text, str) and text.startswith("=")


def test_all_template_sheets_survive(zencoder_output):
    out, _fin, _metrics = zencoder_output
    original = set(load_workbook(find_template()).sheetnames)
    assert original.issubset(set(load_workbook(out).sheetnames))


def test_audit_sheet_records_provenance(zencoder_output):
    out, fin, _metrics = zencoder_output
    ws = load_workbook(out)[AUDIT_SHEET]
    text = "\n".join(
        str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
    assert "Zencoder Inc." in text
    assert "2011" in text
    assert "Audited Financial Statements.pdf" in text
    # Debt was affirmatively absent; OCF was read.
    assert "not on statement (treated as 0)" in text
    assert "Net cash from operating activities" in text
    # The second-pass disagreement is surfaced, not swallowed.
    assert "total revenue: 1 vs 2" in text


def test_audit_sheet_shows_printed_and_scaled_figures(tmp_path):
    fin = ExtractedFinancials.from_dict(
        {"unit_scale": "thousands", "total_revenue": 45000,
         "net_income": 1200, "company_name": "Acme Co"})
    metrics = compute_metrics(fin)
    out = str(tmp_path / "acme.xlsx")
    write_metrics_to_template(render_metrics_for_excel(metrics), out,
                              financials=fin, metrics=metrics)
    ws = load_workbook(out)[AUDIT_SHEET]
    printed = {row[0].value: (row[1].value, row[2].value)
               for row in ws.iter_rows(min_col=1, max_col=3) if row[0].value}
    assert printed["Total revenue"] == (45000, 45000000)


def test_audit_sheet_is_optional(tmp_path):
    out = str(tmp_path / "bare.xlsx")
    write_metrics_to_template([1] * 9, out)
    assert AUDIT_SHEET not in load_workbook(out).sheetnames


def test_regenerating_does_not_duplicate_the_audit_sheet(tmp_path):
    fin = ExtractedFinancials.from_dict(ZENCODER_FY2011)
    out = str(tmp_path / "twice.xlsx")
    for _ in range(2):
        write_metrics_to_template([1] * 9, out, financials=fin)
    names = load_workbook(out).sheetnames
    assert names.count(AUDIT_SHEET) == 1


def test_wrong_number_of_values_is_rejected(tmp_path):
    with pytest.raises(ValueError, match="Expected 9"):
        write_metrics_to_template([1, 2, 3], str(tmp_path / "x.xlsx"))


def test_non_numeric_value_is_rejected(tmp_path):
    with pytest.raises(ValueError, match="not numeric"):
        write_metrics_to_template(["n/a"] + [1] * 8, str(tmp_path / "x.xlsx"))


def test_missing_template_is_reported(tmp_path):
    with pytest.raises(FileNotFoundError):
        write_metrics_to_template([1] * 9, str(tmp_path / "x.xlsx"),
                                  template_path=str(tmp_path / "nope.xlsx"))


def test_output_directory_is_created(tmp_path):
    out = str(tmp_path / "nested" / "deep" / "x.xlsx")
    assert os.path.isfile(write_metrics_to_template([1] * 9, out))


def test_generated_name_is_timestamped():
    from fill_excel import generate_output_name
    name = generate_output_name("/some/dir/Audited Financial Statements.pdf")
    assert name.startswith("Audited Financial Statements_")
    assert name.endswith(".xlsx")
    assert METRIC_ROW == 5  # the contract the template's formulas depend on
